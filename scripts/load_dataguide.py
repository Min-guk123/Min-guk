"""
DataGuide 엑셀(시트1: 종목별 '수정주가'/'시가총액' 시계열, 시트2: 코스닥 '종가지수')을 읽어
엔진 입력(prices, index, universe)으로 바꾼다. 처음 배포할 때 사용자 파일로 초기 데이터를 만들거나,
네이버 수집이 막혔을 때 수동으로 갱신하는 용도.

사용법: python scripts/build.py --source dataguide --file "코스닥 분류.xlsx"
"""
from __future__ import annotations

import openpyxl
import pandas as pd


def _read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr = {str(r[0]).strip(): i for i, r in enumerate(rows[:20]) if r and r[0] is not None}
    code_r, name_r = hdr["Code"], hdr["Name"]
    date_r = next(i for i, r in enumerate(rows[:20]) if r and str(r[0]).replace(" ", "") == "DATE")
    return rows, code_r, name_r, date_r


def load(path: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    stock_ws, index_ws = None, None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=14, max_col=3, values_only=True))
        item = [r for r in rows if r and r[0] == "D A T E"]
        if not item:
            continue
        if item[0][1] in ("수정주가", "시가총액"):
            stock_ws = ws
        elif item[0][1] == "종가지수":
            index_ws = ws
    if stock_ws is None or index_ws is None:
        raise SystemExit("DataGuide 시트(수정주가/시가총액, 종가지수)를 찾지 못했습니다.")

    rows, code_r, name_r, date_r = _read_sheet(stock_ws)
    codes, names, items = rows[code_r], rows[name_r], rows[date_r]
    body = [r for r in rows[date_r + 1:] if r and r[0] is not None]
    dates = pd.to_datetime([r[0] for r in body])
    px, cap, nm = {}, {}, {}
    for c in range(1, len(codes)):
        if not codes[c]:
            continue
        code = str(codes[c]).lstrip("A")
        col = [r[c] if c < len(r) else None for r in body]
        if items[c] == "수정주가":
            px[code] = col
        elif items[c] == "시가총액":
            cap[code] = col
        nm[code] = str(names[c])
    prices = pd.DataFrame(px, index=dates).apply(pd.to_numeric, errors="coerce").sort_index()
    caps = pd.DataFrame(cap, index=dates).apply(pd.to_numeric, errors="coerce").sort_index()
    prices = prices.dropna(how="all")
    caps = caps.loc[prices.index]

    irows, _, _, idate_r = _read_sheet(index_ws)
    ibody = [r for r in irows[idate_r + 1:] if r and r[0] is not None]
    index = pd.Series([r[1] for r in ibody], index=pd.to_datetime([r[0] for r in ibody]))
    index = pd.to_numeric(index, errors="coerce").dropna().sort_index()

    last = caps.ffill().iloc[-1]
    universe = pd.DataFrame({"code": last.index, "name": [nm[c] for c in last.index], "mcap": last.values})
    prices.attrs["caps"] = caps          # 일별 시가총액(원) - 섹터 이평선 위 비율의 시총 가중, 밸류 하단 지표가 쓴다
    return prices, index, universe
